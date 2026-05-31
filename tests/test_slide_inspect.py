"""Tests for the generated-slide consistency checker (slide_inspect).

These power the v4 QA gate: extract a style+geometry signature per text shape, then
assert a 'sibling set' (repeated elements that should look identical) collapses to one
signature. This is what catches the theme-black labels and mixed bullet weights that a
bare 'run has no rPr' grep misses.
"""

from __future__ import annotations

import io

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from src.slide_inspect import consistency_report, slide_signatures


def _styled_label(slide, left, *, text, fill=None, bold=None, font=None, size=None,
                  align=None, anchor=None, width=1.5, height=0.4):
    box = slide.shapes.add_textbox(Inches(left), Inches(1), Inches(width), Inches(height))
    tf = box.text_frame
    if anchor is not None:
        tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    if align is not None:
        p.alignment = align
    run = p.add_run()
    run.text = text
    if bold is not None:
        run.font.bold = bold
    if font is not None:
        run.font.name = font
    if size is not None:
        run.font.size = Pt(size)
    if fill is not None:
        run.font.color.rgb = RGBColor.from_string(fill)
    return box


def _slide_bytes(build):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    build(slide)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_signature_captures_explicit_vs_inherited_style():
    def build(s):
        _styled_label(s, 1, text="A", fill="FFFFFF", bold=True, font="Arial", size=10)
        _styled_label(s, 3, text="B")  # no fill/bold/font -> inherited
    sigs = slide_signatures(_slide_bytes(build))
    assert len(sigs) == 2
    assert sigs[0]["fill"] == "FFFFFF" and sigs[0]["bold"] is True and sigs[0]["font"] == "Arial"
    # the inherited one reads as None on each axis (this is the theme-black trap)
    assert sigs[1]["fill"] is None and sigs[1]["bold"] is None and sigs[1]["font"] is None


def test_report_flags_a_sibling_with_inherited_style():
    def build(s):
        _styled_label(s, 1, text="A", fill="FFFFFF", bold=True, font="Arial", size=10)
        _styled_label(s, 3, text="B", fill="FFFFFF", bold=True, font="Arial", size=10)
        _styled_label(s, 5, text="C")  # the drifted one
    rep = consistency_report(slide_signatures(_slide_bytes(build)), [[0, 1, 2]])
    assert rep["ok"] is False
    g = rep["groups"][0]
    assert g["uniform"] is False
    assert "fill" in g["differing_fields"] and "bold" in g["differing_fields"]


def test_report_uniform_when_siblings_match():
    def build(s):
        for x in (1, 3, 5):
            _styled_label(s, x, text="X", fill="FFFFFF", bold=True, font="Arial", size=10,
                          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM)
    rep = consistency_report(slide_signatures(_slide_bytes(build)), [[0, 1, 2]])
    assert rep["ok"] is True
    assert rep["groups"][0]["uniform"] is True
    assert rep["groups"][0]["differing_fields"] == []


def test_algn_comes_from_the_representative_paragraph_not_a_blank_leading_one():
    # A box whose paragraph[0] is blank (right-aligned) but whose text lives in a
    # center-aligned paragraph[1]: algn must describe the paragraph the run came from.
    def build(s):
        box = s.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
        tf = box.text_frame
        tf.paragraphs[0].alignment = PP_ALIGN.RIGHT  # blank leading paragraph
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.add_run().text = "HELLO"
    sig = slide_signatures(_slide_bytes(build))[0]
    assert sig["text"] == "HELLO"
    assert sig["algn"] == "CENTER"  # not RIGHT (the blank paragraph[0])


def test_report_treats_all_absent_group_as_uniform():
    # A group whose shape indices don't exist (e.g. only non-text/chart shapes, or a
    # typo) must not be reported as width/height drift with empty value lists.
    def build(s):
        _styled_label(s, 1, text="A", fill="FFFFFF", bold=True, font="Arial", size=10)
    rep = consistency_report(slide_signatures(_slide_bytes(build)), [[7, 8]])
    assert rep["ok"] is True
    assert rep["groups"][0]["uniform"] is True
    assert rep["groups"][0]["differing_fields"] == []


def test_report_flags_geometry_drift():
    def build(s):
        _styled_label(s, 1, text="A", fill="FFFFFF", bold=True, font="Arial", size=10, width=1.5)
        _styled_label(s, 3, text="B", fill="FFFFFF", bold=True, font="Arial", size=10, width=2.5)
    rep = consistency_report(slide_signatures(_slide_bytes(build)), [[0, 1]])
    assert rep["ok"] is False
    assert "width" in rep["groups"][0]["differing_fields"]
