"""Tests for serving one reference slide as a standalone .pptx (Corpus.get_slide_pptx)."""

from __future__ import annotations

import base64
import io
import json

from pptx import Presentation
from pptx.util import Inches

from src.poc_corpus import Corpus


def _write_deck(source_dir, filename, titles):
    prs = Presentation()
    blank = prs.slide_layouts[6]
    for t in titles:
        slide = prs.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        box.text_frame.text = t
    prs.save(str(source_dir / filename))


def _corpus(tmp_path, *, source_filename="demo.pptx", write_source=True):
    corpus = tmp_path / "corpus"
    source = tmp_path / "source"
    corpus.mkdir(parents=True)
    source.mkdir(parents=True)
    if write_source:
        _write_deck(source, source_filename, ["SLIDE_A", "SLIDE_B", "SLIDE_C"])
    deck = {
        "source_filename": source_filename,
        "slide_count": 3,
        "design_system": {},
        "slides": [{"index": 0}, {"index": 1}, {"index": 2}],
    }
    (corpus / "demo-deck.tagged.json").write_text(json.dumps(deck), encoding="utf-8")
    return Corpus(corpus, corpus / "assets", source)


def _titles(pptx_bytes):
    prs = Presentation(io.BytesIO(pptx_bytes))
    return [
        s.text_frame.text
        for slide in prs.slides
        for s in slide.shapes
        if s.has_text_frame and s.text_frame.text
    ]


def test_returns_decodable_single_slide_pptx(tmp_path):
    result = _corpus(tmp_path).get_slide_pptx("demo-deck", 1)
    assert result["deck"] == "demo-deck" and result["index"] == 1
    assert result["mime_type"] == (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
    assert result["filename"] == "demo-deck-slide-1.pptx"
    pptx_bytes = base64.b64decode(result["base64"])
    assert _titles(pptx_bytes) == ["SLIDE_B"]


def test_unknown_deck_is_none(tmp_path):
    assert _corpus(tmp_path).get_slide_pptx("nope", 0) is None


def test_out_of_range_index_is_none(tmp_path):
    assert _corpus(tmp_path).get_slide_pptx("demo-deck", 99) is None


def test_missing_source_file_is_none(tmp_path):
    corpus = _corpus(tmp_path, write_source=False)
    assert corpus.get_slide_pptx("demo-deck", 0) is None


def test_resolves_pptx_when_source_filename_names_a_pdf(tmp_path):
    # Corpus JSONs name the .pdf in source_filename; we must find the .pptx sibling.
    corpus = tmp_path / "corpus"
    source = tmp_path / "source"
    corpus.mkdir(parents=True)
    source.mkdir(parents=True)
    _write_deck(source, "demo-deck.pptx", ["ONLY"])
    deck = {"source_filename": "demo-deck.pdf", "slide_count": 1, "slides": [{"index": 0}]}
    (corpus / "demo-deck.tagged.json").write_text(json.dumps(deck), encoding="utf-8")
    result = Corpus(corpus, corpus / "assets", source).get_slide_pptx("demo-deck", 0)
    assert _titles(base64.b64decode(result["base64"])) == ["ONLY"]


def test_resolves_source_by_deck_id_when_filename_absent(tmp_path):
    # deck id is "demo-deck"; if source_filename is missing, fall back to <deck>.pptx
    corpus = tmp_path / "corpus"
    source = tmp_path / "source"
    corpus.mkdir(parents=True)
    source.mkdir(parents=True)
    _write_deck(source, "demo-deck.pptx", ["ONLY"])
    deck = {"slide_count": 1, "slides": [{"index": 0}]}  # no source_filename
    (corpus / "demo-deck.tagged.json").write_text(json.dumps(deck), encoding="utf-8")
    result = Corpus(corpus, corpus / "assets", source).get_slide_pptx("demo-deck", 0)
    assert _titles(base64.b64decode(result["base64"])) == ["ONLY"]
