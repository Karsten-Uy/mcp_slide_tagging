"""A deck is loaded (and visible to every endpoint) only if its source .pptx
exists AND its slide count aligns with the tagged JSON. Others are dropped."""

from __future__ import annotations

import json

from pptx import Presentation
from pptx.util import Inches

from src.poc_corpus import Corpus


def _write_pptx(path, n_slides):
    prs = Presentation()
    blank = prs.slide_layouts[6]
    for i in range(n_slides):
        slide = prs.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        box.text_frame.text = f"S{i}"
    prs.save(str(path))


def _setup(tmp_path, *, json_slides, pptx_slides=None, write_pptx=True):
    """One deck 'deck' with `json_slides` tagged slides; optionally a source .pptx
    with `pptx_slides` (defaults to json_slides) slides."""
    corpus = tmp_path / "corpus"
    source = tmp_path / "source"
    corpus.mkdir(parents=True)
    source.mkdir(parents=True)
    deck = {
        "slide_count": json_slides,
        "design_system": {},
        "slides": [{"index": i, "slide_purpose": "Finding"} for i in range(json_slides)],
    }
    (corpus / "deck.tagged.json").write_text(json.dumps(deck), encoding="utf-8")
    if write_pptx:
        _write_pptx(source / "deck.pptx", json_slides if pptx_slides is None else pptx_slides)
    return Corpus(corpus, corpus / "assets", source)


def test_deck_with_matching_pptx_is_loaded(tmp_path):
    c = _setup(tmp_path, json_slides=3)
    assert "deck" in c.decks
    assert [d["deck"] for d in c.list_decks()] == ["deck"]


def test_audit_reports_loaded_and_excluded_with_reasons(tmp_path):
    corpus = tmp_path / "corpus"
    source = tmp_path / "source"
    corpus.mkdir(parents=True)
    source.mkdir(parents=True)

    def _deck(name, n):
        (corpus / f"{name}.tagged.json").write_text(
            json.dumps({"slide_count": n, "design_system": {}, "slides": [{"index": i} for i in range(n)]}),
            encoding="utf-8",
        )

    _deck("good", 3)
    _write_pptx(source / "good.pptx", 3)        # aligned -> loaded
    _deck("nopptx", 2)                           # no .pptx -> excluded
    _deck("mismatch", 3)
    _write_pptx(source / "mismatch.pptx", 5)     # 3 != 5 -> excluded

    audit = Corpus(corpus, corpus / "assets", source).audit()
    assert audit["loaded"] == ["good"]
    reasons = {e["deck"]: e["reason"] for e in audit["excluded"]}
    assert set(reasons) == {"nopptx", "mismatch"}
    assert "pptx" in reasons["nopptx"].lower()
    assert "mismatch" in reasons["mismatch"].lower() or "!=" in reasons["mismatch"]


def test_deck_without_pptx_is_excluded_everywhere(tmp_path):
    c = _setup(tmp_path, json_slides=3, write_pptx=False)
    assert "deck" not in c.decks
    assert c.list_decks() == []
    assert c.get_deck("deck") is None
    assert c.search_slides() == []
    assert c.corpus_stats()["decks"] == 0


def test_deck_with_mismatched_count_is_excluded(tmp_path):
    c = _setup(tmp_path, json_slides=3, pptx_slides=5)
    assert "deck" not in c.decks
    assert c.get_slide("deck", 0) is None
    assert c.get_slide_pptx("deck", 0) is None
