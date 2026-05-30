"""Corpus.load() uses the precomputed manifest for the admission gate when present,
and falls back to parsing the .pptx only for decks the manifest omits."""

from __future__ import annotations

import json

from pptx import Presentation
from pptx.util import Inches

from src.poc_corpus import Corpus


def _write_pptx(path, n_slides):
    prs = Presentation()
    blank = prs.slide_layouts[6]
    for i in range(n_slides):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = f"S{i}"
    prs.save(str(path))


def _deck_json(corpus, deck_id, n_slides):
    deck = {
        "slide_count": n_slides,
        "design_system": {},
        "slides": [{"index": i} for i in range(n_slides)],
    }
    (corpus / f"{deck_id}.tagged.json").write_text(json.dumps(deck), encoding="utf-8")


def _write_manifest(source, counts):
    (source / "manifest.json").write_text(
        json.dumps({"version": 1, "slide_counts": counts}), encoding="utf-8"
    )


def test_uses_manifest_count_without_parsing_the_pptx(tmp_path):
    # A deliberately unparseable .pptx: if load() admits the deck, it trusted the
    # manifest count and never opened the file.
    corpus = tmp_path / "corpus"
    source = tmp_path / "source"
    corpus.mkdir(parents=True)
    source.mkdir(parents=True)
    _deck_json(corpus, "deck", 3)
    (source / "deck.pptx").write_bytes(b"not a real pptx")
    _write_manifest(source, {"deck.pptx": 3})

    c = Corpus(corpus, corpus / "assets", source)
    assert "deck" in c.decks  # admitted purely from the manifest


def test_manifest_count_mismatch_still_excludes(tmp_path):
    corpus = tmp_path / "corpus"
    source = tmp_path / "source"
    corpus.mkdir(parents=True)
    source.mkdir(parents=True)
    _deck_json(corpus, "deck", 3)
    (source / "deck.pptx").write_bytes(b"not a real pptx")
    _write_manifest(source, {"deck.pptx": 5})  # manifest says 5, tags say 3

    c = Corpus(corpus, corpus / "assets", source)
    assert "deck" not in c.decks


def test_falls_back_to_parsing_when_manifest_omits_deck(tmp_path):
    # No manifest entry for this deck -> the real .pptx is parsed (must be valid).
    corpus = tmp_path / "corpus"
    source = tmp_path / "source"
    corpus.mkdir(parents=True)
    source.mkdir(parents=True)
    _deck_json(corpus, "deck", 2)
    _write_pptx(source / "deck.pptx", 2)
    _write_manifest(source, {})  # manifest present but empty

    c = Corpus(corpus, corpus / "assets", source)
    assert "deck" in c.decks


def test_no_manifest_file_still_parses(tmp_path):
    # Backward compatible: no manifest.json at all -> parse as before.
    corpus = tmp_path / "corpus"
    source = tmp_path / "source"
    corpus.mkdir(parents=True)
    source.mkdir(parents=True)
    _deck_json(corpus, "deck", 2)
    _write_pptx(source / "deck.pptx", 2)

    c = Corpus(corpus, corpus / "assets", source)
    assert "deck" in c.decks
