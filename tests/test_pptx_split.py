"""Tests for extracting one slide from a deck into a standalone .pptx."""

from __future__ import annotations

import io

import pytest
from pptx import Presentation
from pptx.util import Inches

from src.pptx_split import extract_single_slide


def _deck_bytes(titles: list[str]) -> bytes:
    """Build a multi-slide .pptx with one title textbox per slide."""
    prs = Presentation()
    blank = prs.slide_layouts[6]  # blank layout
    for t in titles:
        slide = prs.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        box.text_frame.text = t
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _titles(pptx_bytes: bytes) -> list[str]:
    prs = Presentation(io.BytesIO(pptx_bytes))
    out: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text:
                out.append(shape.text_frame.text)
    return out


def test_extracts_only_the_requested_slide():
    deck = _deck_bytes(["SLIDE_A", "SLIDE_B", "SLIDE_C"])
    result = extract_single_slide(deck, 1)
    prs = Presentation(io.BytesIO(result))
    assert len(prs.slides) == 1
    assert _titles(result) == ["SLIDE_B"]


def test_first_slide_index_zero():
    deck = _deck_bytes(["SLIDE_A", "SLIDE_B", "SLIDE_C"])
    assert _titles(extract_single_slide(deck, 0)) == ["SLIDE_A"]


def test_last_slide():
    deck = _deck_bytes(["SLIDE_A", "SLIDE_B", "SLIDE_C"])
    assert _titles(extract_single_slide(deck, 2)) == ["SLIDE_C"]


def test_result_is_valid_reopenable_pptx():
    deck = _deck_bytes(["ONLY"])
    result = extract_single_slide(deck, 0)
    # Reopening without raising is the validity check.
    assert len(Presentation(io.BytesIO(result)).slides) == 1


def test_index_out_of_range_raises():
    deck = _deck_bytes(["SLIDE_A", "SLIDE_B"])
    with pytest.raises(IndexError):
        extract_single_slide(deck, 5)


def test_negative_index_raises():
    deck = _deck_bytes(["SLIDE_A", "SLIDE_B"])
    with pytest.raises(IndexError):
        extract_single_slide(deck, -1)
