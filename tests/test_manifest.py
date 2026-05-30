"""Tests for the precomputed slide-count manifest (build / write / load)."""

from __future__ import annotations

import json

from pptx import Presentation
from pptx.util import Inches

from src.manifest import build_manifest, load_manifest, write_manifest


def _write_pptx(path, n_slides):
    prs = Presentation()
    blank = prs.slide_layouts[6]
    for i in range(n_slides):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text_frame.text = f"S{i}"
    prs.save(str(path))


def test_build_manifest_counts_every_pptx(tmp_path):
    _write_pptx(tmp_path / "alpha.pptx", 2)
    _write_pptx(tmp_path / "beta.pptx", 5)
    m = build_manifest(tmp_path)
    assert m["slide_counts"] == {"alpha.pptx": 2, "beta.pptx": 5}
    assert m["version"] >= 1


def test_build_manifest_ignores_non_pptx(tmp_path):
    _write_pptx(tmp_path / "alpha.pptx", 1)
    (tmp_path / "notes.pdf").write_bytes(b"x")
    (tmp_path / "manifest.json").write_text("{}", encoding="utf-8")
    assert set(build_manifest(tmp_path)["slide_counts"]) == {"alpha.pptx"}


def test_write_then_load_roundtrips(tmp_path):
    _write_pptx(tmp_path / "alpha.pptx", 3)
    out = tmp_path / "manifest.json"
    n = write_manifest(tmp_path, out)
    assert n == 1 and out.exists()
    assert load_manifest(out) == {"alpha.pptx": 3}


def test_load_missing_manifest_is_empty(tmp_path):
    assert load_manifest(tmp_path / "nope.json") == {}
